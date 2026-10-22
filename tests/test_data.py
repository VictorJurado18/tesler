"""Parquet / JSON / CSV / XLSX through Arrow: convert, peek, combine."""
import json
from pathlib import Path

import pytest

pa = pytest.importorskip("pyarrow")
import pyarrow.parquet as pq  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from dragshift import convert, data  # noqa: E402
from dragshift.util import UserError  # noqa: E402


def br_csv(path: Path) -> Path:
    path.write_text("id;cliente;cep;valor;pago\n1;Ana;01310-100;1.234,56;sim\n2;Bia;00123;7,5;não\n3;Caio;;;sim\n",
                    encoding="utf-8")
    return path


def test_csv_to_parquet_types(tmp_path):
    out = convert.convert(br_csv(tmp_path / "v.csv"), "sheet", "PARQUET")[0]
    t = pq.read_table(out)
    assert out.suffix == ".parquet"
    assert str(t.schema.field("id").type) == "int64"
    assert str(t.schema.field("valor").type) == "double"
    assert str(t.schema.field("cep").type) == "string"          # codes stay text
    assert t.column("valor").to_pylist() == [1234.56, 7.5, None]
    assert t.column("cep").to_pylist() == ["01310-100", "00123", None]


def test_parquet_to_csv_xlsx_json(tmp_path):
    src = tmp_path / "d.parquet"
    pq.write_table(pa.table({"nome": ["Ana", "Bia"], "valor": [10.5, 3.0], "qtd": [1, None]}), src)
    csv_out = convert.convert(src, "sheet", "CSV")[0]
    assert csv_out.read_text(encoding="utf-8-sig").splitlines() == ["nome;valor;qtd", "Ana;10,5;1", "Bia;3,0;"] \
        or csv_out.read_text(encoding="utf-8-sig").splitlines() == ["nome;valor;qtd", "Ana;10,5;1", "Bia;3;"]
    ws = load_workbook(convert.convert(src, "sheet", "XLSX")[0]).active
    assert [[c.value for c in r] for r in ws.iter_rows()] == [["nome", "valor", "qtd"], ["Ana", 10.5, 1], ["Bia", 3.0, None]]
    js = json.loads(convert.convert(src, "sheet", "JSON")[0].read_text(encoding="utf-8"))
    assert js == [{"nome": "Ana", "valor": 10.5, "qtd": 1}, {"nome": "Bia", "valor": 3.0, "qtd": None}]


def test_nested_json_and_ndjson(tmp_path):
    src = tmp_path / "api.json"
    src.write_text(json.dumps({"page": 1, "items": [
        {"id": 1, "user": {"name": "Ana", "city": "SP"}, "tags": ["a", "b"]},
        {"id": 2, "user": {"name": "Bia"}, "extra": True}]}), encoding="utf-8")
    t = data.load_table(src)
    assert t.column_names == ["id", "user.name", "user.city", "tags", "extra"]
    assert t.column("tags").to_pylist() == ['["a", "b"]', None]
    nd = tmp_path / "log.jsonl"
    nd.write_text('{"a": 1}\n{"a": 2, "b": "x"}\n', encoding="utf-8")
    assert data.load_table(nd).column("b").to_pylist() == [None, "x"]
    bad = tmp_path / "bad.json"
    bad.write_text("{oops", encoding="utf-8")
    with pytest.raises(UserError):
        data.load_table(bad)


def test_peek(tmp_path):
    info = data.peek(br_csv(tmp_path / "v.csv"))
    assert (info["rows"], info["cols"]) == (3, 5)
    by = {c["name"]: c for c in info["columns"]}
    assert by["valor"]["type"] == "decimal" and by["valor"]["nulls"] == 1
    assert (by["id"]["min"], by["id"]["max"]) == (1, 3)
    assert by["pago"]["distinct"] == 2
    assert info["preview"][0][:2] == [1, "Ana"]


def test_combine_mixed_files(tmp_path):
    a = br_csv(tmp_path / "jan.csv")
    b = tmp_path / "fev.xlsx"
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["id", "cliente", "valor", "obs"])
    wb.active.append([4, "Duda", 99, "novo"])
    wb.save(b)
    out = data.combine_files([a, b])[0]
    assert out.name == "jan Combined.xlsx"
    rows = [[c.value for c in r] for r in load_workbook(out).active.iter_rows()]
    assert rows[0] == ["id", "cliente", "cep", "valor", "pago", "obs", "arquivo_origem"]
    assert rows[-1] == [4, "Duda", None, 99, None, "novo", "fev.xlsx"]
    assert len(rows) == 5


def test_combine_csvs_stays_csv(tmp_path):
    out = data.combine_files([br_csv(tmp_path / "a.csv"), br_csv(tmp_path / "b.csv")])[0]
    assert out.suffix == ".csv" and len(out.read_text(encoding="utf-8-sig").splitlines()) == 7
