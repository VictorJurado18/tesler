"""Tabular data files: Parquet, JSON / JSON Lines, CSV/TSV and Excel, through Arrow tables.

load_table() reads any of them into a pyarrow.Table (column types inferred; nested JSON
flattened to "a.b" columns), write_table() writes one out. Peek (stats per column) and
Combine (stack several files, columns matched by name) build on the same two functions.
"""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from .util import UserError, sibling

PARQUET_EXTS = {".parquet", ".pq"}
JSON_EXTS = {".json", ".jsonl", ".ndjson"}
CSV_EXTS = {".csv", ".tsv"}
XLSX_EXTS = {".xlsx", ".xlsm"}
EXCEL_MAX_ROWS = 1_048_575


def pa():
    try:
        import pyarrow
        return pyarrow
    except Exception as e:  # noqa: BLE001
        raise UserError("The data component isn't installed yet. Open Tesler with "
                        "“Iniciar Tesler.bat” once to install it.") from e


def kind(path: str | Path) -> str:
    ext = Path(path).suffix.lower()
    if ext in PARQUET_EXTS:
        return "parquet"
    if ext in JSON_EXTS:
        return "json"
    if ext in XLSX_EXTS:
        return "xlsx"
    return "csv"


def targets(src_ext: str) -> list[str]:
    k = kind("x" + src_ext)
    return {"parquet": ["CSV", "XLSX", "JSON"], "json": ["CSV", "XLSX", "PARQUET"],
            "csv": ["XLSX", "PDF", "PARQUET", "JSON"], "xlsx": ["CSV", "PDF", "PARQUET", "JSON"]}[k]


# ---------------------------------------------------------------- building tables
def _column(values: list):
    """Arrow array with a sensible type for a Python column (mixed → string)."""
    p = pa()
    kinds = {type(v) for v in values if v is not None and v != ""}
    vals = [None if v == "" else v for v in values]
    if not kinds:
        return p.array([None] * len(values), p.string())
    if kinds <= {bool}:
        return p.array(vals, p.bool_())
    if kinds <= {int}:
        if all(v is None or -2**63 <= v < 2**63 for v in vals):
            return p.array(vals, p.int64())
    if kinds <= {int, float}:
        return p.array([None if v is None else float(v) for v in vals], p.float64())
    if kinds <= {dt.datetime, dt.date}:
        if kinds == {dt.date}:
            return p.array(vals, p.date32())
        return p.array([v if isinstance(v, dt.datetime) or v is None else dt.datetime(v.year, v.month, v.day)
                        for v in vals], p.timestamp("us"))
    return p.array([None if v is None else (v if isinstance(v, str) else _text(v)) for v in vals], p.string())


def _text(v) -> str:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return str(v)


def _unique_names(names: list) -> list[str]:
    out, seen = [], {}
    for i, n in enumerate(names):
        n = str(n).strip() if n not in (None, "") else f"coluna_{i + 1}"
        if n in seen:
            seen[n] += 1
            n = f"{n}_{seen[n]}"
        else:
            seen[n] = 1
        out.append(n)
    return out


def table_from_rows(header: list, rows: list[list]):
    names = _unique_names(header)
    width = len(names)
    cols = [[] for _ in range(width)]
    for r in rows:
        r = list(r)[:width] + [None] * (width - len(r))
        for i, v in enumerate(r):
            cols[i].append(v)
    return pa().table({n: _column(c) for n, c in zip(names, cols)})


def _flatten(rec, prefix: str = "", out: dict | None = None) -> dict:
    out = {} if out is None else out
    if isinstance(rec, dict):
        for k, v in rec.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, dict):
                _flatten(v, key, out)
            else:
                out[key] = v
    else:
        out[prefix or "valor"] = rec
    return out


def _json_records(src: Path) -> list:
    text = src.read_text(encoding="utf-8-sig")
    stripped = text.lstrip()
    if src.suffix.lower() in (".jsonl", ".ndjson") or (stripped.startswith("{") and "\n{" in stripped):
        try:
            return [json.loads(ln) for ln in text.splitlines() if ln.strip()]
        except json.JSONDecodeError:
            pass
    try:
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise UserError(f"Invalid JSON (line {e.lineno}).") from e
    if isinstance(data, list):
        return data
    if isinstance(data, dict):           # {"data": [...]} / {"items": [...]}: use the biggest list of records
        lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if lists:
            return max(lists, key=len)
        return [data]
    raise UserError("This JSON isn't a table (expected a list of records).")


# ---------------------------------------------------------------- load / write
def load_table(src: str | Path):
    src = Path(src)
    k = kind(src)
    p = pa()
    if k == "parquet":
        import pyarrow.parquet as pq
        try:
            return pq.read_table(src)
        except Exception as e:  # noqa: BLE001
            raise UserError("This Parquet file couldn't be read.") from e
    if k == "json":
        recs = [_flatten(r) for r in _json_records(src)]
        if not recs:
            raise UserError("This file is empty.")
        names: dict[str, None] = {}
        for r in recs:
            for key in r:
                names.setdefault(key, None)
        header = list(names)
        return table_from_rows(header, [[r.get(n) for n in header] for r in recs])
    if k == "xlsx":
        from .sheets import read_xlsx
        sheets = [(n, r) for n, r in read_xlsx(src) if r]
        if not sheets:
            raise UserError("This spreadsheet is empty.")
        rows = sheets[0][1]
        return table_from_rows(rows[0], rows[1:])
    from .sheets import read_csv, to_number, uses_comma_decimal
    rows, delim = read_csv(src)
    if not rows:
        raise UserError("This file is empty.")
    comma = uses_comma_decimal(rows, delim)
    del p
    return table_from_rows(rows[0], [[to_number(c, comma) for c in r] for r in rows[1:]])


def _py_rows(table):
    cols = [c.to_pylist() for c in table.columns]
    return [list(r) for r in zip(*cols)] if cols else []


def write_table(table, out: Path, fmt: str) -> Path:
    fmt = fmt.upper()
    if fmt == "PARQUET":
        import pyarrow.parquet as pq
        pq.write_table(table, out, compression="snappy")
    elif fmt == "JSON":
        names = table.column_names
        recs = [dict(zip(names, r)) for r in _py_rows(table)]
        out.write_text(json.dumps(recs, ensure_ascii=False, indent=1, default=_text), encoding="utf-8")
    elif fmt == "CSV":
        import csv
        from .sheets import fmt_cell
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(table.column_names)
            for r in _py_rows(table):
                w.writerow([fmt_cell(v) if not isinstance(v, (dict, list)) else _text(v) for v in r])
    elif fmt == "XLSX":
        if table.num_rows > EXCEL_MAX_ROWS:
            raise UserError(f"Too many rows for Excel ({table.num_rows:,}). Convert to CSV or Parquet instead."
                            .replace(",", "."))
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        wb = Workbook(write_only=False)
        ws = wb.active
        ws.append(table.column_names)
        for r in _py_rows(table):
            ws.append([_text(v) if isinstance(v, (dict, list)) else v for v in r])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, name in enumerate(table.column_names, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, len(str(name)) + 4), 40)
        wb.save(out)
    else:
        raise ValueError(fmt)
    return out


def convert_data(src: str | Path, target: str) -> list[Path]:
    target = target.upper()
    ext = {"PARQUET": ".parquet", "JSON": ".json", "CSV": ".csv", "XLSX": ".xlsx"}[target]
    return [write_table(load_table(src), sibling(src, ext=ext), target)]


# ---------------------------------------------------------------- peek
def peek(src: str | Path, preview_rows: int = 200) -> dict:
    """{'rows', 'cols', 'columns': [{name, type, nulls, distinct, min, max, example}], 'preview': [[...]], 'header'}"""
    import pyarrow.compute as pc
    t = load_table(src)
    cols = []
    for name, col in zip(t.column_names, t.columns):
        typ = col.type
        info = {"name": name, "type": friendly_type(typ), "nulls": col.null_count, "distinct": None,
                "min": None, "max": None, "example": None}
        try:
            info["distinct"] = pc.count_distinct(col).as_py()
        except Exception:
            pass
        p = pa()
        if p.types.is_integer(typ) or p.types.is_floating(typ) or p.types.is_temporal(typ):
            try:
                mm = pc.min_max(col).as_py()
                info["min"], info["max"] = mm["min"], mm["max"]
            except Exception:
                pass
        non_null = col.drop_null()
        if len(non_null):
            info["example"] = non_null[0].as_py()
        cols.append(info)
    return {"rows": t.num_rows, "cols": t.num_columns, "columns": cols, "header": t.column_names,
            "preview": _py_rows(t.slice(0, preview_rows))}


def friendly_type(typ) -> str:
    p = pa()
    if p.types.is_integer(typ):
        return "integer"
    if p.types.is_floating(typ) or p.types.is_decimal(typ):
        return "decimal"
    if p.types.is_boolean(typ):
        return "boolean"
    if p.types.is_timestamp(typ):
        return "date/time"
    if p.types.is_date(typ):
        return "date"
    if p.types.is_string(typ) or p.types.is_large_string(typ):
        return "text"
    if p.types.is_null(typ):
        return "empty"
    return str(typ)


# ---------------------------------------------------------------- combine
def combine(paths: list[str | Path], source_column: str = "arquivo_origem"):
    """Stack several files; columns matched by name (missing → empty); a column with
    the source file name is added. Conflicting types become text (int+float → decimal)."""
    p = pa()
    tables = [(Path(s), load_table(s)) for s in paths]
    order: dict[str, None] = {}
    types: dict[str, set] = {}
    for _, t in tables:
        for f in t.schema:
            order.setdefault(f.name, None)
            if not p.types.is_null(f.type):
                types.setdefault(f.name, set()).add(f.type)
    target: dict[str, object] = {}
    for name in order:
        ts = types.get(name, set())
        if len(ts) == 1:
            target[name] = next(iter(ts))
        elif ts and all(p.types.is_integer(x) or p.types.is_floating(x) for x in ts):
            target[name] = p.float64()
        else:
            target[name] = p.string()
    parts = []
    for path, t in tables:
        cols = {}
        for name, typ in target.items():
            if name in t.column_names:
                c = t.column(name)
                if c.type != typ:
                    c = c.cast(typ) if not p.types.is_string(typ) else p.array(
                        [None if v is None else _text(v) for v in c.to_pylist()], p.string())
                cols[name] = c
            else:
                cols[name] = p.nulls(t.num_rows, typ)
        cols[source_column] = p.array([path.name] * t.num_rows, p.string())
        parts.append(p.table(cols))
    return p.concat_tables(parts)


def combine_files(paths: list[str | Path]) -> list[Path]:
    table = combine(paths)
    kinds = {kind(s) for s in paths}
    if kinds == {"parquet"}:
        fmt = "PARQUET"
    elif kinds == {"json"}:
        fmt = "JSON"
    elif kinds == {"csv"} or table.num_rows > EXCEL_MAX_ROWS:
        fmt = "CSV"
    else:
        fmt = "XLSX"
    ext = {"PARQUET": ".parquet", "JSON": ".json", "CSV": ".csv", "XLSX": ".xlsx"}[fmt]
    return [write_table(table, sibling(paths[0], "Combined", ext), fmt)]
