"""Spreadsheets: CSV/TSV ⇄ XLSX, and either → PDF.

Brazilian CSVs usually use ';' and decimal commas ("1.234,56"); that is detected so numbers
arrive in Excel as numbers. Codes with leading zeros (CEP, CPF, product codes) stay text.
"""
from __future__ import annotations

import csv
import html
import io
import re
import sys
import traceback
from pathlib import Path

from .convert import fitz, read_text
from .util import UserError, sibling

SHEET_TARGETS_CSV = ["XLSX", "PDF"]
SHEET_TARGETS_XLSX = ["CSV", "PDF"]

_INT = re.compile(r"^-?\d+$")
_DOT_DEC = re.compile(r"^-?\d+\.\d+$")
_BR_NUM = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+,\d+$")


# ---------------------------------------------------------------- reading
def sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample[:20000], delimiters=";,\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
        return max(counts, key=counts.get) if any(counts.values()) else ","


def read_csv(src: str | Path) -> tuple[list[list[str]], str]:
    text = read_text(src)
    delim = "\t" if Path(src).suffix.lower() == ".tsv" else sniff_delimiter(text)
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    return rows, delim


def to_number(s: str, comma_decimal: bool):
    """'12' → 12, '3.5' → 3.5, '1.234,56' → 1234.56 (when the file uses decimal commas).
    Anything else (incl. leading zeros like '00123') stays a string."""
    t = s.strip()
    if not t or len(t) > 18:
        return s
    if (t.startswith("0") and len(t) > 1 and not t.startswith(("0,", "0."))) or \
       (t.startswith("-0") and len(t) > 2 and not t.startswith(("-0,", "-0."))):
        return s                                   # leading zero → code, keep as text
    if _INT.match(t):
        return int(t)
    if comma_decimal:
        if _BR_NUM.match(t):
            return float(t.replace(".", "").replace(",", "."))
        return s
    if _DOT_DEC.match(t):
        return float(t)
    return s


def uses_comma_decimal(rows: list[list[str]], delim: str) -> bool:
    if delim == ",":
        return False
    cells = [c.strip() for r in rows[:200] for c in r]
    comma = sum(1 for c in cells if re.match(r"^-?[\d.]+,\d+$", c))
    dot = sum(1 for c in cells if _DOT_DEC.match(c))
    return comma > dot


def read_xlsx(src: str | Path) -> list[tuple[str, list[list]]]:
    """[(sheet name, rows)] with computed values (as last saved by Excel)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise UserError("This spreadsheet couldn't be opened (is it password-protected?).") from e
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None or str(v).strip() == "" for v in rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(r) if v not in (None, "")), default=0) for r in rows),
                    default=0)
        sheets.append((ws.title, [r[:width] for r in rows]))
    wb.close()
    return sheets


def fmt_cell(v, comma_decimal: bool = True) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, float):
        s = f"{v:.10g}" if abs(v) < 1e15 else str(v)
        return s.replace(".", ",") if comma_decimal else s
    if hasattr(v, "strftime"):
        return v.strftime("%d/%m/%Y %H:%M") if getattr(v, "hour", 0) or getattr(v, "minute", 0) \
            else v.strftime("%d/%m/%Y")
    return str(v)


# ---------------------------------------------------------------- conversions
def csv_to_xlsx(src: str | Path) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    rows, delim = read_csv(src)
    if not rows:
        raise UserError("This file is empty.")
    comma = uses_comma_decimal(rows, delim)
    wb = Workbook()
    ws = wb.active
    ws.title = Path(src).stem[:31] or "Sheet1"
    for i, r in enumerate(rows):
        ws.append(r if i == 0 else [to_number(c, comma) for c in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i in range(1, max(len(r) for r in rows) + 1):   # rough auto-width
        width = max((len(str(r[i - 1])) for r in rows[:500] if len(r) >= i), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60)
    out = sibling(src, ext=".xlsx")
    wb.save(out)
    return [out]


def xlsx_to_csv(src: str | Path) -> list[Path]:
    """One CSV per sheet, ';' + UTF-8 with BOM: opens correctly in Excel PT-BR."""
    sheets = [(n, r) for n, r in read_xlsx(src) if r]
    if not sheets:
        raise UserError("This spreadsheet is empty.")
    outs = []
    for name, rows in sheets:
        label = name if len(sheets) > 1 else ""
        out = sibling(src, label, ".csv")
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            for r in rows:
                w.writerow([fmt_cell(v) for v in r])
        outs.append(out)
    return outs


TABLE_CSS = """
* { font-family: sans-serif; }
h2 { font-size: 12pt; margin: 0 0 6pt 0; }
table { border-collapse: collapse; width: 100%; }
td, th { border: 0.5px solid #999; padding: 2pt 4pt; font-size: 8pt; vertical-align: top; }
th { background-color: #e8eef7; font-weight: bold; }
td.n { text-align: right; }
"""


def _table_html(title: str | None, rows: list[list], header: bool = True) -> str:
    parts = [f"<h2>{html.escape(title)}</h2>"] if title else []
    parts.append("<table>")
    for i, r in enumerate(rows):
        tag = "th" if header and i == 0 else "td"
        cells = []
        for v in r:
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            cls = ' class="n"' if num and tag == "td" else ""
            cells.append(f"<{tag}{cls}>{html.escape(fmt_cell(v))}</{tag}>")
        parts.append("<tr>" + "".join(cells) + "</tr>")
    parts.append("</table>")
    return "".join(parts)


def _render_tables(sections: list[tuple[str | None, list[list]]], out: Path) -> None:
    widest = max((max((len(r) for r in rows), default=0) for _, rows in sections), default=0)
    page = fitz.paper_rect("a4-l" if widest > 6 else "a4")
    writer = fitz.DocumentWriter(str(out))
    for title, rows in sections:          # each sheet starts on a new page
        story = fitz.Story(html=_table_html(title, rows), user_css=TABLE_CSS)
        where = page + (28, 28, -28, -28)
        more = True
        while more:
            dev = writer.begin_page(page)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
    writer.close()


def _excel_to_pdf(src: Path, out: Path) -> bool:
    """Microsoft Excel (if installed) gives the exact print layout."""
    if sys.platform != "win32":
        return False
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except Exception:
        return False
    pythoncom.CoInitialize()
    xl = None
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(str(src.resolve()), ReadOnly=True)
        try:
            wb.ExportAsFixedFormat(0, str(out.resolve()))      # xlTypePDF
        finally:
            wb.Close(False)
        return out.exists()
    except Exception:
        traceback.print_exc()
        return False
    finally:
        if xl is not None:
            from .util import quiet_faults
            with quiet_faults():
                try:
                    xl.Quit()
                except Exception:
                    pass
        pythoncom.CoUninitialize()


def sheet_to_pdf(src: str | Path) -> list[Path]:
    src = Path(src)
    out = sibling(src, ext=".pdf")
    if src.suffix.lower() in (".xlsx", ".xlsm"):
        if _excel_to_pdf(src, out):
            return [out]
        sheets = [(n, r) for n, r in read_xlsx(src) if r]
        sections = [(n if len(sheets) > 1 else None, r) for n, r in sheets]
    else:
        rows, delim = read_csv(src)
        comma = uses_comma_decimal(rows, delim)
        sections = [(None, [rows[0]] + [[to_number(c, comma) for c in r] for r in rows[1:]])] if rows else []
    if not sections:
        raise UserError("This spreadsheet is empty.")
    _render_tables(sections, out)
    return [out]


def convert_sheet(src: str | Path, target: str) -> list[Path]:
    from . import data
    target = target.upper()
    k = data.kind(src)
    if k in ("parquet", "json") or target in ("PARQUET", "JSON"):
        return data.convert_data(src, target)
    if target == "PDF":
        return sheet_to_pdf(src)
    if target == "XLSX":
        return csv_to_xlsx(src)
    if target == "CSV":
        return xlsx_to_csv(src)
    raise ValueError(f"Can't convert spreadsheet to {target}")


def targets(src_ext: str) -> list[str]:
    return SHEET_TARGETS_XLSX if src_ext.lower() in (".xlsx", ".xlsm") else SHEET_TARGETS_CSV


# ---------------------------------------------------------------- PDF tables → XLSX
def _clean_cell(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def pdf_tables(src: str | Path) -> list[tuple[int, list[list[str]]]]:
    """[(page number 1-based, rows)] for every table PyMuPDF detects (lines or text alignment)."""
    from .convert import open_pdf
    found = []
    with open_pdf(src) as doc:
        for pg in doc:
            tabs = pg.find_tables()                       # ruled tables
            if not tabs.tables:
                tabs = pg.find_tables(strategy="text")    # borderless (aligned columns)
            for t in tabs.tables:
                # read each cell by its box: Table.extract() can scramble characters
                # ("1.250,00" → "125000 . ,") with some fonts
                rows = [[_clean_cell(pg.get_text("text", clip=c)) if c else "" for c in row.cells]
                        for row in t.rows]
                rows = [r for r in rows if any(r)]
                if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                    found.append((pg.number + 1, rows))
    return found


def pdf_to_xlsx(src: str | Path) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    tables = pdf_tables(src)
    if not tables:
        from .convert import open_pdf
        with open_pdf(src) as doc:
            scanned = sum(len(p.get_text().strip()) for p in doc) < 20
        raise UserError("No tables found. This looks like a scanned PDF: run OCR first."
                        if scanned else "No tables found in this PDF.")
    wb = Workbook()
    wb.remove(wb.active)
    per_page: dict[int, int] = {}
    for page, rows in tables:
        per_page[page] = per_page.get(page, 0) + 1
        ws = wb.create_sheet(f"p{page}" + (f" ({per_page[page]})" if per_page[page] > 1 else ""))
        comma = uses_comma_decimal(rows, ";")
        for i, r in enumerate(rows):
            ws.append(r if i == 0 else [to_number(c, comma) for c in r])
        for c in ws[1]:
            c.font = Font(bold=True)
        for i in range(1, max(len(r) for r in rows) + 1):
            width = max((len(str(r[i - 1])) for r in rows if len(r) >= i), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 50)
    out = sibling(src, ext=".xlsx")
    wb.save(out)
    return [out]
