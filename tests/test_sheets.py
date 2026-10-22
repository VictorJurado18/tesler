"""CSV/XLSX conversions."""
import csv
from pathlib import Path

import fitz
from openpyxl import Workbook, load_workbook

from dragshift import convert, sheets
from dragshift.util import category


def test_category_and_targets():
    assert category("a.csv") == category("b.xlsx") == "sheet"
    assert convert.targets_for("sheet", ".csv") == ["XLSX", "PDF"]
    assert convert.targets_for("sheet", ".xlsx") == ["CSV", "PDF"]


def test_brazilian_csv_to_xlsx(tmp_path: Path):
    src = tmp_path / "vendas.csv"
    src.write_bytes("Produto;CEP;Qtd;Preço\nCafé;01310-100;3;1.234,56\nPão;00123;10;2,5\n".encode("cp1252"))
    out = convert.convert(src, "sheet", "XLSX")[0]
    ws = load_workbook(out).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["Produto", "CEP", "Qtd", "Preço"]
    assert rows[1] == ["Café", "01310-100", 3, 1234.56]
    assert rows[2] == ["Pão", "00123", 10, 2.5]          # leading zero stays text


def test_us_csv_numbers(tmp_path: Path):
    src = tmp_path / "us.csv"
    src.write_text("a,b\n1,2.5\n3,x\n", encoding="utf-8")
    ws = load_workbook(convert.convert(src, "sheet", "XLSX")[0]).active
    assert [c.value for c in ws[2]] == [1, 2.5]


def test_xlsx_to_csv_one_per_sheet(tmp_path: Path):
    src = tmp_path / "rel.xlsx"
    wb = Workbook()
    wb.active.title = "Jan"
    wb.active.append(["Nome", "Valor"])
    wb.active.append(["Ana", 10.5])
    ws2 = wb.create_sheet("Fev")
    ws2.append(["Nome", "Valor"])
    ws2.append(["Bia", 7])
    wb.save(src)
    outs = convert.convert(src, "sheet", "CSV")
    assert [o.name for o in outs] == ["rel Jan.csv", "rel Fev.csv"]
    rows = list(csv.reader(outs[0].read_text(encoding="utf-8-sig").splitlines(), delimiter=";"))
    assert rows == [["Nome", "Valor"], ["Ana", "10,5"]]


def test_sheet_to_pdf(tmp_path: Path):
    src = tmp_path / "t.csv"
    src.write_text("Nome;Valor\n" + "".join(f"Item {i};{i},5\n" for i in range(120)), encoding="utf-8")
    out = convert.convert(src, "sheet", "PDF")[0]
    d = fitz.open(out)
    assert d.page_count >= 2 and "Item 119" in "".join(p.get_text() for p in d)


def _statement_pdf(path: Path) -> Path:
    doc = fitz.open()
    pg = doc.new_page()
    rows = [("Data", "Descrição", "Valor (R$)", "Saldo (R$)"), ("01/10/2026", "Saldo anterior", "", "1.250,00"),
            ("05/10/2026", "Boleto energia", "-189,47", "1.060,53")]
    x, y0 = [50, 140, 360, 460, 560], 90
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            pg.insert_text((x[j] + 4, y0 + i * 22 + 15), c, fontsize=10)
    for i in range(len(rows) + 1):
        pg.draw_line((x[0], y0 + i * 22), (x[-1], y0 + i * 22))
    for xx in x:
        pg.draw_line((xx, y0), (xx, y0 + len(rows) * 22))
    pg2 = doc.new_page()                                   # borderless, aligned columns
    for i, r in enumerate([("Código", "Produto", "Qtd"), ("00123", "Café 500g", "3"), ("00456", "Açúcar 1kg", "2")]):
        for j, c in enumerate(r):
            pg2.insert_text(([60, 160, 360][j], 100 + i * 20), c, fontsize=11)
    doc.save(path)
    return path


def test_pdf_tables_to_xlsx(tmp_path: Path):
    out = convert.convert(_statement_pdf(tmp_path / "extrato.pdf"), "pdf", "XLSX")[0]
    wb = load_workbook(out)
    assert wb.sheetnames == ["p1", "p2"]
    p1 = [[c.value for c in r] for r in wb["p1"].iter_rows()]
    assert p1[1] == ["01/10/2026", "Saldo anterior", None, 1250]
    assert p1[2] == ["05/10/2026", "Boleto energia", -189.47, 1060.53]
    p2 = [[c.value for c in r] for r in wb["p2"].iter_rows()]
    assert p2[1] == ["00123", "Café 500g", 3]


def test_pdf_without_tables(tmp_path: Path):
    import pytest
    from dragshift.util import UserError
    p = tmp_path / "t.pdf"
    d = fitz.open()
    d.new_page().insert_text((72, 72), "Só um parágrafo de texto, nada de tabela aqui.")
    d.save(p)
    with pytest.raises(UserError):
        convert.convert(p, "pdf", "XLSX")
